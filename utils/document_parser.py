"""Text extraction helpers for uploaded resume documents."""

import re

from PyPDF2 import PdfReader
from docx import Document


class DocumentParsingError(Exception):
    """Raised when an uploaded resume cannot be read."""


def clean_text(text):
    """Remove excessive whitespace while retaining useful line and paragraph breaks."""
    if not text:
        return ""
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    lines = [re.sub(r"[ \t]+", " ", line).strip() for line in text.split("\n")]
    cleaned_lines, previous_was_blank = [], False
    for line in lines:
        if not line:
            if not previous_was_blank and cleaned_lines:
                cleaned_lines.append("")
            previous_was_blank = True
        else:
            cleaned_lines.append(line)
            previous_was_blank = False
    return "\n".join(cleaned_lines).strip()


def extract_text_from_pdf(file):
    """Extract and clean text from a PDF file-like object."""
    try:
        file.seek(0)
        reader = PdfReader(file)
        text = "\n\n".join(page.extract_text() or "" for page in reader.pages)
    except Exception as exc:
        raise DocumentParsingError("We couldn't read this PDF. It may be corrupted, password-protected, or image-only.") from exc
    return clean_text(text)


def extract_text_from_docx(file):
    """Extract and clean text from a DOCX file-like object."""
    try:
        file.seek(0)
        document = Document(file)
        parts = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text.strip() for cell in row.cells if cell.text.strip()))
    except Exception as exc:
        raise DocumentParsingError("We couldn't read this DOCX file. Please upload a valid, non-corrupted Word document.") from exc
    return clean_text("\n".join(parts))


def extract_resume_text(file):
    """Choose the appropriate extractor for an uploaded PDF or DOCX resume."""
    filename = (getattr(file, "name", "") or "").lower()
    if filename.endswith(".pdf"):
        return extract_text_from_pdf(file)
    if filename.endswith(".docx"):
        return extract_text_from_docx(file)
    raise DocumentParsingError("Unsupported file type. Please upload a PDF or DOCX resume.")
